import sys
import os
import threading
import pymysql
import openpyxl
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import ttkbootstrap as tb
from ttkbootstrap.constants import *

# 如果不需要PIL，可以只使用iconbitmap，但为了跨平台保留PIL
try:
    from PIL import Image, ImageTk
except ImportError:
    Image = None


class IPCheckerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("IP配置比对工具")
        self.root.geometry("1000x650")

        # ---------- 设置窗口图标（强制刷新） ----------
        self.set_window_icon()
        self.root.update_idletasks()  # 确保图标立即更新

        # ---------- 数据库连接 ----------
        db_frame = tb.LabelFrame(root, text="数据库连接")
        db_frame.pack(fill="x", padx=10, pady=5)

        # 预设IP列表
        self.ip_list = [
            "192.168.1.213", "192.168.1.216", "192.168.1.219",
            "192.168.1.237", "192.168.1.238", "192.168.1.249",
            "192.168.1.243", "192.168.1.128", "192.168.1.210",
            "192.168.1.152", "192.168.1.153"
        ]

        tb.Label(db_frame, text="主机:").grid(row=0, column=0, padx=5, pady=2, sticky="w")
        self.host_combo = ttk.Combobox(db_frame, values=self.ip_list, width=15)
        self.host_combo.grid(row=0, column=1, padx=5, pady=2)
        self.host_combo.set(self.ip_list[0])

        tb.Label(db_frame, text="端口:").grid(row=0, column=2, padx=5, pady=2, sticky="w")
        self.port_entry = tb.Entry(db_frame, width=8)
        self.port_entry.grid(row=0, column=3, padx=5, pady=2)
        self.port_entry.insert(0, "3306")

        tb.Label(db_frame, text="用户名:").grid(row=0, column=4, padx=5, pady=2, sticky="w")
        self.user_entry = tb.Entry(db_frame, width=12)
        self.user_entry.grid(row=0, column=5, padx=5, pady=2)
        self.user_entry.insert(0, "root")

        tb.Label(db_frame, text="密码:").grid(row=0, column=6, padx=5, pady=2, sticky="w")
        self.pwd_entry = tb.Entry(db_frame, width=12, show="")
        self.pwd_entry.grid(row=0, column=7, padx=5, pady=2)
        self.pwd_entry.insert(0, "mysql")

        tb.Label(db_frame, text="数据库:").grid(row=0, column=8, padx=5, pady=2, sticky="w")
        self.db_entry = tb.Entry(db_frame, width=12)
        self.db_entry.grid(row=0, column=9, padx=5, pady=2)
        self.db_entry.insert(0, "base")

        # ---------- Excel文件选择 ----------
        file_frame = tb.LabelFrame(root, text="Excel文件")
        file_frame.pack(fill="x", padx=10, pady=5)

        self.file_path_var = tk.StringVar(value="未选择文件")
        tb.Label(file_frame, textvariable=self.file_path_var).pack(side="left", padx=5)
        tb.Button(file_frame, text="选择Excel文件", command=self.select_file, bootstyle="primary").pack(side="left", padx=5)

        # ---------- 比对按钮 ----------
        self.compare_btn = tb.Button(root, text="开始比对", command=self.start_compare, bootstyle="success", state="disabled")
        self.compare_btn.pack(pady=10)

        # ---------- 结果表格 ----------
        tree_frame = tb.Frame(root)
        tree_frame.pack(fill="both", expand=True, padx=10, pady=5)

        columns = ("系统编号", "站点编号", "Excel IP", "数据库IP列表（站点）", "匹配")
        self.tree = tb.Treeview(tree_frame, columns=columns, show="headings", bootstyle="primary")
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=150, anchor="center")
        self.tree.column("数据库IP列表（站点）", width=250)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        self.tree.pack(side="left", fill="both", expand=True)

        # ---------- 状态栏 ----------
        self.status_var = tk.StringVar(value="就绪")
        status_bar = tb.Label(root, textvariable=self.status_var, bootstyle="secondary", anchor="w")
        status_bar.pack(fill="x", padx=10, pady=5)

        self.excel_path = None
        self.thread = None

    # ---------- 图标设置（强制使用.ico文件） ----------
    def set_window_icon(self):
        """直接使用 .ico 文件设置图标，Windows 下最可靠"""
        icon_path = "ip_check_icon.ico"
        if not os.path.exists(icon_path):
            print(f"警告: 图标文件 {icon_path} 未找到，将使用默认图标")
            return

        try:
            # Windows 下使用 iconbitmap 设置标题栏和任务栏图标
            self.root.iconbitmap(default=icon_path)
            # 同时调用 iconphoto 以支持其他平台（需要 Pillow）
            if Image:
                img = Image.open(icon_path)
                photo = ImageTk.PhotoImage(img)
                self.icon_img = photo  # 防止被垃圾回收
                self.root.iconphoto(True, photo)
            # 强制刷新
            self.root.update()
        except Exception as e:
            print(f"设置图标失败: {e}")

    # ---------- 其余方法（select_file, start_compare 等）不变 ----------
    def select_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel文件", "*.xlsx *.xls")])
        if path:
            self.excel_path = path
            self.file_path_var.set(path)
            self.compare_btn.config(state="normal")

    def start_compare(self):
        if not self.excel_path:
            messagebox.showwarning("警告", "请先选择Excel文件")
            return

        db_config = {
            "host": self.host_combo.get().strip(),
            "port": int(self.port_entry.get().strip()),
            "user": self.user_entry.get().strip(),
            "password": self.pwd_entry.get().strip(),
            "database": self.db_entry.get().strip()
        }
        if not all([db_config["host"], db_config["user"], db_config["database"]]):
            messagebox.showwarning("警告", "请填写完整的数据库连接信息")
            return

        for item in self.tree.get_children():
            self.tree.delete(item)

        self.compare_btn.config(state="disabled")
        self.status_var.set("正在比对，请稍候...")

        self.thread = threading.Thread(target=self._compare_task, args=(db_config,), daemon=True)
        self.thread.start()
        self._check_thread()

    def _compare_task(self, db_config):
        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
            if "IP配置" not in wb.sheetnames:
                raise ValueError("Excel 文件缺少名为 'IP配置' 的工作表")
            sheet = wb["IP配置"]

            headers = [cell.value for cell in sheet[1]]
            required_cols = ["系统编号", "站点编号", "IP", "PORT"]
            col_indices = {}
            for col_name in required_cols:
                if col_name not in headers:
                    raise ValueError(f"工作表缺少必需的列：'{col_name}'")
                col_indices[col_name] = headers.index(col_name)

            excel_rows = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                excel_rows.append({
                    "系统编号": row[col_indices["系统编号"]],
                    "站点编号": row[col_indices["站点编号"]],
                    "IP": row[col_indices["IP"]],
                    "PORT": row[col_indices["PORT"]]
                })

            if not excel_rows:
                raise ValueError("Excel 工作表中没有数据")

            conn = pymysql.connect(
                host=db_config["host"],
                port=db_config["port"],
                user=db_config["user"],
                password=db_config["password"],
                database=db_config["database"],
                charset="utf8mb4"
            )
            cursor = conn.cursor()
            cursor.execute("SELECT sys_no, station_no, IP FROM dic_mss_netaddress_config")
            rows = cursor.fetchall()
            cursor.close()
            conn.close()

            ip_map = {}
            for sys_no, station_no, ip in rows:
                key = (sys_no, station_no)
                ip_map.setdefault(key, []).append(ip)

            result_list = []
            for row in excel_rows:
                sys_no = row["系统编号"]
                station_no = row["站点编号"]
                excel_ip = row["IP"]

                key = (sys_no, station_no)
                db_ips = ip_map.get(key, [])
                match = excel_ip in db_ips
                db_ip_display = ", ".join(db_ips) if db_ips else "(无)"

                result_list.append({
                    "系统编号": sys_no,
                    "站点编号": station_no,
                    "Excel IP": excel_ip,
                    "数据库IP列表": db_ip_display,
                    "匹配": "是" if match else "否"
                })

            result_list.sort(key=lambda x: (x["系统编号"], x["站点编号"]))

            self._result_list = result_list
            self._error_msg = None

        except pymysql.err.OperationalError as e:
            if "Unknown column 'station_no'" in str(e):
                self._error_msg = "数据库表缺少 'station_no' 字段，请检查字段名或修改代码。"
            else:
                self._error_msg = f"数据库错误：{e}"
            self._result_list = None
        except Exception as e:
            self._error_msg = str(e)
            self._result_list = None

        self.root.after(0, self._update_ui)

    def _check_thread(self):
        if self.thread and self.thread.is_alive():
            self.root.after(500, self._check_thread)

    def _update_ui(self):
        self.compare_btn.config(state="normal")
        if hasattr(self, "_error_msg") and self._error_msg:
            self.status_var.set("比对出错")
            messagebox.showerror("错误", f"比对失败：{self._error_msg}")
            return

        result_list = getattr(self, "_result_list", [])
        if not result_list:
            self.status_var.set("没有数据")
            return

        for row_data in result_list:
            values = (
                row_data["系统编号"],
                row_data["站点编号"],
                row_data["Excel IP"],
                row_data["数据库IP列表"],
                row_data["匹配"]
            )
            item = self.tree.insert("", "end", values=values)
            if row_data["匹配"] == "否":
                self.tree.tag_configure("unmatch", background="#ffcccc")
                self.tree.item(item, tags=("unmatch",))

        total = len(result_list)
        match_count = sum(1 for r in result_list if r["匹配"] == "是")
        self.status_var.set(f"比对完成：总计 {total} 条，匹配 {match_count} 条，不匹配 {total - match_count} 条")
        messagebox.showinfo("比对完成", f"总计 {total} 条记录，匹配 {match_count} 条，不匹配 {total - match_count} 条。")


if __name__ == "__main__":
    root = tb.Window(themename="superhero")
    app = IPCheckerApp(root)
    root.mainloop()