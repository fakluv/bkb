import socket
import time
import struct
import threading
from datetime import datetime
import tkinter as tk
from tkinter import scrolledtext, messagebox, filedialog
import os

# 尝试导入 openpyxl，若失败则提示并退出
try:
    from openpyxl import load_workbook
except ImportError:
    messagebox.showerror("缺少依赖", "请先安装 openpyxl 库：pip install openpyxl")
    raise SystemExit

# 全局停止标志
stop_sending = False

def build_fixed_template():
    """构建固定部分的报文（bytearray），动态字段留空。总长99字节（紧急制动回路状态仅1字节）"""
    msg = bytearray(99)

    msg[0] = 0xC9
    msg[1], msg[2] = 0x20, 0x0B
    # 3~6 序列号，7~10 时间戳，发送时动态填充
    msg[11], msg[12] = 0x14, 0x01
    msg[13] = 0x11
    msg[14], msg[15] = 0xFF, 0xFF        # ATP 24V电源板状态
    msg[16] = 0xFF                       # ATP主机板状态
    msg[17:22] = bytes([0xFF] * 5)       # ATP输入板状态
    msg[22:25] = bytes([0xFF] * 3)       # ATP输出板1状态
    msg[25:28] = bytes([0xFF] * 3)       # ATP输出板2状态
    msg[28] = 0xFF                       # ATP公共板状态及校轮状态
    msg[29] = 0xFF                       # 紧急制动回路状态（仅1字节）
    msg[30] = 0xFF                       # ATO 24v电源板状态
    msg[31], msg[32] = 0xFF, 0xFF        # ATO输入板状态
    msg[33:37] = bytes([0xFF] * 4)       # ATO输出板状态
    msg[37] = 0xFF                       # ATO主机板状态
    msg[38] = 0xC3                       # BTM总体状态
    msg[39:47] = bytes([0xFF] * 8)       # 1端BTM各板卡状态信息（8字节）
    msg[47:55] = bytes([0xFF] * 8)       # 2端BTM各板卡状态信息（8字节）
    msg[55], msg[56] = 0x8F, 0x8F        # 1端速传状态, 2端速传状态
    msg[57] = 0xDF                       # 辅助测速传感器状态
    msg[58:64] = bytes([0x02, 0x01, 0x05, 0x04, 0x09, 0x08])  # 日检状态
    msg[64], msg[65] = 0x00, 0x00        # 司机号
    msg[66] = 0x03                       # 当前控制级别
    msg[67] = 0x02                       # 当前驾驶模式
    msg[68], msg[69] = 0x00, 0x0B        # 站台号
    msg[70] = 0xAA                       # 停车精度标志
    msg[71], msg[72] = 0x00, 0x66        # 停车精度数值
    # 应答器编号 73~76 动态填充
    msg[77] = 0x00                       # 通信中断的受控联锁ID
    msg[78] = 0x00                       # 通信中断的受控ZC ID
    msg[79] = 0x00                       # 通信中断的受控ATS ID
    msg[80:84] = bytes([0xFF] * 4)       # 控制端VOBC内网通信状态等4个状态字节
    msg[84:90] = bytes([0x00] * 6)       # 列车紧急制动原因
    msg[90:92] = bytes([0x00] * 2)       # 列车切断牵引原因
    msg[92:94] = bytes([0x00] * 2)       # 列车丢失位置原因
    msg[94] = 0x00                       # ATB失败原因
    msg[95], msg[96] = 0x00, 0x00        # 打开车门不关闭
    msg[97], msg[98] = 0x00, 0x00        # 辅助测速传感器状态

    return msg

def packet_to_formatted_string(packet):
    """将报文转换为指定的多行字符串格式，用于显示（总长99字节）"""
    lines = []
    lines.append(f"{packet[0]:02X} #msg_id#")
    lines.append(f"{packet[1]:02X} {packet[2]:02X} #版本号#")

    seq = struct.unpack_from('>I', packet, 3)[0]
    lines.append(f"[{packet[3]:02X} {packet[4]:02X} {packet[5]:02X} {packet[6]:02X}] #数据来源端ATP的序列号#")

    timestamp = struct.unpack_from('>I', packet, 7)[0]
    dt_str = datetime.fromtimestamp(timestamp).strftime('%Y-%m-%d--%H:%M:%S') if timestamp else "无效时间"
    lines.append(f"[{packet[7]:02X} {packet[8]:02X} {packet[9]:02X} {packet[10]:02X}] #数据来源一端ATP时间 ({dt_str})#")

    lines.append(f"{packet[11]:02X} {packet[12]:02X} #消息源设备ID#")
    lines.append(f"{packet[13]:02X} #主控端&消息来源...")
    lines.append(f"{packet[14]:02X} {packet[15]:02X} #ATP 24V电源板状态#")
    lines.append(f"{packet[16]:02X} #ATP主机板状态#")
    lines.append(f"{packet[17]:02X} {packet[18]:02X} {packet[19]:02X} {packet[20]:02X} {packet[21]:02X} #ATP输入板状态#")
    lines.append(f"{packet[22]:02X} {packet[23]:02X} {packet[24]:02X} #ATP输出板1状态#")
    lines.append(f"{packet[25]:02X} {packet[26]:02X} {packet[27]:02X} #ATP输出板2状态#")
    lines.append(f"{packet[28]:02X} #ATP公共板状态及校轮状态#")
    lines.append(f"{packet[29]:02X} #紧急制动回路状态#")
    lines.append(f"{packet[30]:02X} #ATO 24v电源板状态#")
    lines.append(f"{packet[31]:02X} {packet[32]:02X} #ATO输入板状态#")
    lines.append(f"{packet[33]:02X} {packet[34]:02X} {packet[35]:02X} {packet[36]:02X} #ATO输出板状态#")
    lines.append(f"{packet[37]:02X} #ATO主机板状态#")
    lines.append(f"{packet[38]:02X} #BTM总体状态#")
    lines.append(f"{packet[39]:02X} {packet[40]:02X} {packet[41]:02X} {packet[42]:02X} {packet[43]:02X} {packet[44]:02X} {packet[45]:02X} {packet[46]:02X} #1端BTM各板卡状态信息#")
    lines.append(f"{packet[47]:02X} {packet[48]:02X} {packet[49]:02X} {packet[50]:02X} {packet[51]:02X} {packet[52]:02X} {packet[53]:02X} {packet[54]:02X} #2端BTM各板卡状态信息#")
    lines.append(f"{packet[55]:02X} #1端速传状态#")
    lines.append(f"{packet[56]:02X} #2端速传状态#")
    lines.append(f"{packet[57]:02X} #辅助测速传感器状态 Dx加速度计 Ex:雷达 #")
    lines.append(f"{packet[58]:02X} {packet[59]:02X} {packet[60]:02X} {packet[61]:02X} {packet[62]:02X} {packet[63]:02X} #日检状态#")
    lines.append(f"{packet[64]:02X} {packet[65]:02X} #司机号...")
    lines.append(f"{packet[66]:02X} #当前控制级别...")
    lines.append(f"{packet[67]:02X} #当前驾驶模式...")
    lines.append(f"{packet[68]:02X} {packet[69]:02X} #站台号...")
    lines.append(f"{packet[70]:02X} #停车精度标志...")
    lines.append(f"{packet[71]:02X} {packet[72]:02X} #停车精度数值...")
    lines.append(f"{packet[73]:02X} {packet[74]:02X} {packet[75]:02X} {packet[76]:02X} #低2个字节是应答器编号#")
    lines.append(f"{packet[77]:02X} # 通信中断的受控联锁ID#")
    lines.append(f"{packet[78]:02X} #通信中断的受控ZC ID#")
    lines.append(f"{packet[79]:02X} #通信中断的受控ATS ID#")
    lines.append(f"{packet[80]:02X} #控制端VOBC内网通信状态#")
    lines.append(f"{packet[81]:02X} #VOBC自关断报警#")
    lines.append(f"{packet[82]:02X} #XX车XX端牵引反馈异常/制动反馈异常#")
    lines.append(f"{packet[83]:02X} #客室门屏蔽门状态#")
    lines.append(f"{packet[84]:02X} {packet[85]:02X} {packet[86]:02X} {packet[87]:02X} {packet[88]:02X} {packet[89]:02X} #列车紧急制动原因...")
    lines.append(f"{packet[90]:02X} {packet[91]:02X} #列车切断牵引原因...")
    lines.append(f"{packet[92]:02X} {packet[93]:02X} #列车丢失位置原因...")
    lines.append(f"{packet[94]:02X} # ATB失败原因#")
    lines.append(f"{packet[95]:02X} {packet[96]:02X} #打开车门不关闭#")
    lines.append(f"{packet[97]:02X} {packet[98]:02X} #辅助测速传感器状态#")

    return "\n".join(lines)

def read_balise_numbers(filepath):
    """
    从电子地图文件中读取应答器表索引编号
    返回：整数编号列表，若读取失败返回 None
    """
    try:
        wb = load_workbook(filepath, data_only=True)
        if "应答器表" not in wb.sheetnames:
            messagebox.showerror("错误", "文件中未找到名为“应答器表”的工作表")
            return None
        ws = wb["应答器表"]

        # 查找第4行的“索引编号”列
        header_row = 4
        col_idx = None
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=header_row, column=col).value
            if cell_value and str(cell_value).strip() == "索引编号":
                col_idx = col
                break
        if col_idx is None:
            messagebox.showerror("错误", "在第4行未找到列名“索引编号”")
            return None

        # 从第5行开始读取该列的数字
        numbers = []
        for row in range(5, ws.max_row + 1):
            val = ws.cell(row=row, column=col_idx).value
            if val is None:
                continue
            try:
                num = int(val)
                numbers.append(num)
            except (ValueError, TypeError):
                # 非数字内容自动跳过
                continue
        return numbers

    except Exception as e:
        messagebox.showerror("读取文件失败", f"读取Excel文件时出错：{str(e)}")
        return None

def send_worker(local_ip, target_ip, target_port, balise_list, print_enabled, text_widget):
    """在后台线程中发送报文（每个编号连续发送5秒）"""
    global stop_sending
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.bind((local_ip, 0))
    except Exception as e:
        text_widget.insert(tk.END, f"绑定本机IP失败: {e}\n")
        text_widget.see(tk.END)
        return

    total_numbers = len(balise_list)
    sent = 0
    for balise in balise_list:
        if stop_sending:
            break
        # 同一应答器编号连续发送5秒，每秒1包
        for i in range(5):
            if stop_sending:
                break

            packet = build_fixed_template()
            now = int(time.time())
            struct.pack_into('>I', packet, 3, now)            # 序列号
            struct.pack_into('>I', packet, 7, now)            # 时间戳
            struct.pack_into('>I', packet, 73, balise & 0xFFFF)  # 应答器编号（偏移73）

            try:
                sock.sendto(packet, (target_ip, target_port))
            except Exception as e:
                text_widget.insert(tk.END, f"发送失败: {e}\n")
                text_widget.see(tk.END)
                stop_sending = True
                break

            # 根据复选框决定是否打印详细报文
            if print_enabled.get():
                formatted = packet_to_formatted_string(packet)
                header = f"\n========== 发送包 #{balise} (第{i+1}/5次, 时间戳={now}) ==========\n"
                footer = "\n=====================================\n"
                text_widget.insert(tk.END, header + formatted + footer)
            else:
                text_widget.insert(tk.END, f"发送包 #{balise} (第{i+1}/5次, 时间戳={now}) 成功\n")
            text_widget.see(tk.END)

            sent += 1

            # 如果不是最后一个包，间隔1秒
            if not stop_sending and not (balise == balise_list[-1] and i == 4):
                time.sleep(1)

        if stop_sending:
            break

    sock.close()
    if stop_sending:
        text_widget.insert(tk.END, f"\n发送已停止，共发送 {sent} 个报文。\n")
    else:
        text_widget.insert(tk.END, f"\n发送完成！共发送 {sent} 个报文。\n")
    text_widget.see(tk.END)
    root.after(0, lambda: btn_start.config(state=tk.NORMAL))
    root.after(0, lambda: btn_stop.config(state=tk.DISABLED))

def start_sending():
    global stop_sending
    filepath = entry_file.get().strip()
    if not filepath:
        messagebox.showerror("错误", "请选择电子地图文件")
        return
    if not os.path.isfile(filepath):
        messagebox.showerror("错误", "文件不存在，请重新选择")
        return

    # 读取编号列表
    balise_list = read_balise_numbers(filepath)
    if balise_list is None:
        return
    if not balise_list:
        messagebox.showerror("错误", "未能从文件中读取到任何有效的数字编号")
        return

    # 显示读取结果
    text_output.insert(tk.END, f"\n从文件中读取到 {len(balise_list)} 个应答器编号：{balise_list[:10]}{'...' if len(balise_list)>10 else ''}\n\n")
    text_output.see(tk.END)

    local_ip = entry_local_ip.get().strip()
    target_ip = entry_target_ip.get().strip()
    try:
        port = int(entry_port.get().strip())
    except:
        messagebox.showerror("错误", "端口必须是整数")
        return

    stop_sending = False
    btn_start.config(state=tk.DISABLED)
    btn_stop.config(state=tk.NORMAL)

    thread = threading.Thread(target=send_worker,
                              args=(local_ip, target_ip, port, balise_list, print_var, text_output),
                              daemon=True)
    thread.start()

def stop_sending_func():
    global stop_sending
    stop_sending = True
    btn_stop.config(state=tk.DISABLED)

def clear_output():
    text_output.delete(1.0, tk.END)

def browse_file():
    filename = filedialog.askopenfilename(
        title="选择电子地图文件",
        filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
    )
    if filename:
        entry_file.delete(0, tk.END)
        entry_file.insert(0, filename)
        # 自动预览读取到的编号
        balise_list = read_balise_numbers(filename)
        if balise_list:
            text_output.delete(1.0, tk.END)
            text_output.insert(tk.END, f"文件加载成功，共 {len(balise_list)} 个有效编号：\n{balise_list}\n\n")
        else:
            text_output.delete(1.0, tk.END)
            text_output.insert(tk.END, "文件读取失败或未找到有效编号，请检查工作表格式。\n")

# 创建主窗口
root = tk.Tk()
root.title("CBTC 应答器编号 UDP 发送器（从电子地图读取）")
root.geometry("900x700")

# 输入区域框架
frame_input = tk.Frame(root)
frame_input.pack(pady=10, padx=10, fill=tk.X)

# 本机IP
tk.Label(frame_input, text="本机IP:").grid(row=0, column=0, sticky=tk.W)
entry_local_ip = tk.Entry(frame_input, width=15)
entry_local_ip.insert(0, "192.168.1.51")
entry_local_ip.grid(row=0, column=1, padx=5)

# 目标IP
tk.Label(frame_input, text="目标IP:").grid(row=0, column=2, sticky=tk.W)
entry_target_ip = tk.Entry(frame_input, width=15)
entry_target_ip.insert(0, "192.168.1.238")
entry_target_ip.grid(row=0, column=3, padx=5)

# 端口
tk.Label(frame_input, text="端口:").grid(row=0, column=4, sticky=tk.W)
entry_port = tk.Entry(frame_input, width=8)
entry_port.insert(0, "10330")
entry_port.grid(row=0, column=5, padx=5)

# 电子地图文件
tk.Label(frame_input, text="电子地图文件:").grid(row=1, column=0, sticky=tk.W, pady=5)
entry_file = tk.Entry(frame_input, width=50)
entry_file.grid(row=1, column=1, columnspan=3, padx=5, sticky=tk.W)
btn_browse = tk.Button(frame_input, text="浏览", command=browse_file, width=8)
btn_browse.grid(row=1, column=4, padx=5, sticky=tk.W)

# 打印报文复选框
print_var = tk.BooleanVar(value=True)
chk_print = tk.Checkbutton(frame_input, text="打印报文详情", variable=print_var)
chk_print.grid(row=1, column=5, sticky=tk.W, padx=10)

# 按钮框架
frame_buttons = tk.Frame(root)
frame_buttons.pack(pady=5)

btn_start = tk.Button(frame_buttons, text="开始发送", command=start_sending, width=12)
btn_start.pack(side=tk.LEFT, padx=10)

btn_stop = tk.Button(frame_buttons, text="停止发送", command=stop_sending_func, state=tk.DISABLED, width=12)
btn_stop.pack(side=tk.LEFT, padx=10)

btn_clear = tk.Button(frame_buttons, text="清空输出", command=clear_output, width=12)
btn_clear.pack(side=tk.LEFT, padx=10)

# 输出区域（带滚动条）
frame_output = tk.Frame(root)
frame_output.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
text_output = scrolledtext.ScrolledText(frame_output, wrap=tk.WORD, font=("Courier New", 9))
text_output.pack(fill=tk.BOTH, expand=True)

root.mainloop()